"""CSV / XLSX import and XLSX export.

Import is two-step so nothing is overwritten by accident: the upload is
parsed into a preview (stored server-side under a random token, owned by
the uploading user), and only an explicit confirm creates a new spreadsheet
or adds the data as NEW sheets of an existing one.
"""
import csv
import io
import json
import os
import re
import secrets
import time
from datetime import date, datetime, time as dtime
from typing import Optional

from flask import current_app

from app.sheets import validators as V

MAX_IMPORT_ROWS = 50000   # same as the most rows a sheet can have
MAX_IMPORT_CELLS = 500000
PREVIEW_ROWS = 15
TOKEN_TTL_SECONDS = 2 * 3600
_TOKEN_RE = re.compile(r'^[A-Za-z0-9_-]{20,64}$')


class ImportError_(ValueError):
    """User-facing import problem."""


# ── Parsing ───────────────────────────────────────────────────────────────

def _to_text(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == 0:
            return value.strftime('%Y-%m-%d')
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, dtime):
        return value.strftime('%H:%M:%S')
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    text = str(value)
    return text if text != '' else None


def _parse_csv(data: bytes, filename: str) -> list:
    for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:20000], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i >= MAX_IMPORT_ROWS:
            raise ImportError_(f'The file has more than {MAX_IMPORT_ROWS:,} rows.')
        rows.append(row[:V.MAX_COLS])
    name = os.path.splitext(os.path.basename(filename))[0][:V.MAX_SHEET_NAME] or 'Imported'
    cells = [(r, c, v, None) for r, row in enumerate(rows) for c, v in enumerate(row) if v != '']
    if len(cells) > MAX_IMPORT_CELLS:
        raise ImportError_(f'The file has more than {MAX_IMPORT_CELLS:,} filled cells.')
    return [{'name': name, 'cells': cells, 'props': {}, 'frozen_rows': 0, 'frozen_cols': 0}]


def _argb(color) -> Optional[str]:
    try:
        rgb = color.rgb if color is not None else None
    except Exception:
        return None
    if isinstance(rgb, str) and re.match(r'^[0-9A-Fa-f]{8}$', rgb):
        return '#' + rgb[2:].lower()
    return None


def _xlsx_fmt(cell) -> Optional[dict]:
    fmt = {}
    font = cell.font
    if font is not None:
        if font.b:
            fmt['b'] = True
        if font.i:
            fmt['i'] = True
        if font.u:
            fmt['u'] = True
        if font.strike:
            fmt['s'] = True
        if font.sz and int(font.sz) != 11:
            fmt['fs'] = int(font.sz)
        c = _argb(font.color)
        if c and c != '#000000':
            fmt['c'] = c
    fill = cell.fill
    if fill is not None and fill.fill_type == 'solid':
        bg = _argb(fill.fgColor)
        if bg and bg != '#ffffff':
            fmt['bg'] = bg
    al = cell.alignment
    if al is not None:
        if al.horizontal in ('left', 'center', 'right'):
            fmt['ha'] = al.horizontal
        if al.vertical in ('top', 'center', 'bottom'):
            fmt['va'] = 'middle' if al.vertical == 'center' else al.vertical
        if al.wrap_text:
            fmt['wrap'] = 'wrap'
    nf = (cell.number_format or 'General')
    if nf != 'General':
        if '%' in nf:
            fmt['nf'] = {'t': 'percent', 'd': nf.split('.')[-1].count('0') if '.' in nf else 0}
        elif any(sym in nf for sym in ('$', '€', '£', 'Rs', '₨')):
            cur = next(sym for sym in ('Rs', '₨', '$', '€', '£') if sym in nf)
            fmt['nf'] = {'t': 'currency', 'cur': cur, 'd': 2 if '.00' in nf else 0}
        elif re.search(r'[dy]', nf, re.I) and re.search(r'h', nf, re.I):
            fmt['nf'] = {'t': 'datetime'}
        elif re.search(r'[dy]', nf, re.I):
            fmt['nf'] = {'t': 'date'}
        elif re.search(r'h', nf, re.I):
            fmt['nf'] = {'t': 'time'}
        elif '0.0' in nf:
            fmt['nf'] = {'t': 'decimal', 'd': nf.split('.')[-1].count('0')}
        elif '#,##0' in nf:
            fmt['nf'] = {'t': 'number', 'd': 0}
    for side, key in (('top', 'bt'), ('bottom', 'bb'), ('left', 'bl'), ('right', 'br')):
        b = getattr(cell.border, side, None) if cell.border is not None else None
        if b is not None and b.style:
            fmt[key] = {'w': 2 if b.style in ('medium', 'thick') else 1,
                        's': {'dashed': 'dashed', 'dotted': 'dotted', 'double': 'double'}.get(b.style, 'solid'),
                        'c': _argb(b.color) or '#000000'}
    return V.clean_fmt(fmt)


def _parse_xlsx(data: bytes) -> list:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import column_index_from_string
        wb = load_workbook(io.BytesIO(data), data_only=False)
    except Exception:
        raise ImportError_('This file could not be read as an Excel workbook (.xlsx).')
    sheets, total = [], 0
    for ws in wb.worksheets[:50]:
        cells = []
        if ws.max_row > MAX_IMPORT_ROWS:
            raise ImportError_(f'Sheet "{ws.title}" has more than {MAX_IMPORT_ROWS:,} rows.')
        for row in ws.iter_rows(max_col=min(ws.max_column, V.MAX_COLS)):
            for cell in row:
                value = cell.value
                fmt = _xlsx_fmt(cell) if cell.has_style else None
                if value is None and not fmt:
                    continue
                text = _to_text(value)
                if isinstance(value, str) and value.startswith('=') and cell.data_type != 'f':
                    text = "'" + value
                cells.append((cell.row - 1, cell.column - 1, text, fmt))
        total += len(cells)
        if total > MAX_IMPORT_CELLS:
            raise ImportError_(f'The workbook has more than {MAX_IMPORT_CELLS:,} filled cells.')
        props = {'cw': {}, 'rh': {}, 'merges': [], 'hc': [], 'hr': []}
        for letter, dim in ws.column_dimensions.items():
            try:
                idx = column_index_from_string(letter) - 1
            except ValueError:
                continue
            if dim.hidden:
                props['hc'].append(idx)
            if dim.width:
                props['cw'][str(idx)] = max(20, min(2000, int(dim.width * 7 + 5)))
        for idx, dim in ws.row_dimensions.items():
            if dim.hidden:
                props['hr'].append(idx - 1)
            if dim.height:
                props['rh'][str(idx - 1)] = max(12, min(1000, int(dim.height / 0.75)))
        for rng in ws.merged_cells.ranges:
            props['merges'].append({'r1': rng.min_row - 1, 'c1': rng.min_col - 1,
                                    'r2': rng.max_row - 1, 'c2': rng.max_col - 1})
        frozen_rows = frozen_cols = 0
        if ws.freeze_panes:
            m = re.match(r'^([A-Z]+)(\d+)$', str(ws.freeze_panes))
            if m:
                frozen_cols = column_index_from_string(m.group(1)) - 1
                frozen_rows = int(m.group(2)) - 1
        sheets.append({'name': ws.title[:V.MAX_SHEET_NAME] or 'Sheet', 'cells': cells,
                       'props': props, 'frozen_rows': min(frozen_rows, 50), 'frozen_cols': min(frozen_cols, 26),
                       'hidden': ws.sheet_state != 'visible',
                       'color': _argb(ws.sheet_properties.tabColor) if ws.sheet_properties.tabColor else None})
    if not sheets:
        raise ImportError_('The workbook has no sheets.')
    return sheets


def parse_upload(filename: str, data: bytes) -> list:
    ext = os.path.splitext(filename or '')[1].lower()
    if ext in ('.csv', '.tsv', '.txt'):
        return _parse_csv(data, filename)
    if ext in ('.xlsx', '.xlsm'):
        return _parse_xlsx(data)
    if ext == '.xls':
        raise ImportError_('Old .xls files are not supported - save it as .xlsx or .csv first.')
    raise ImportError_('Choose a .csv or .xlsx file.')


def _column_kind(values: list) -> str:
    seen = [v for v in values if v not in (None, '')]
    if not seen:
        return 'empty'
    if all(re.match(r'^-?[\d,]*\.?\d+%?$', v.strip()) for v in seen):
        return 'number'
    if all(re.match(r'^\d{4}-\d{2}-\d{2}', v.strip()) or re.match(r'^\d{1,2}/\d{1,2}/\d{2,4}$', v.strip())
           for v in seen):
        return 'date'
    if all(v.startswith('=') for v in seen):
        return 'formula'
    return 'text'


def preview(sheets: list) -> list:
    out = []
    for s in sheets:
        grid = {}
        max_r = max_c = -1
        for r, c, v, _ in s['cells']:
            max_r, max_c = max(max_r, r), max(max_c, c)
            if r < PREVIEW_ROWS + 1:
                grid[(r, c)] = v
        cols = max_c + 1
        rows = [[grid.get((r, c)) or '' for c in range(cols)] for r in range(min(max_r + 1, PREVIEW_ROWS + 1))]
        body_values = {c: [] for c in range(cols)}
        for r, c, v, _ in s['cells']:
            if 0 < r <= 200 and v is not None:
                body_values[c].append(v)
        header = rows[0] if rows else []
        out.append({
            'name': s['name'], 'rows': rows, 'row_count': max_r + 1, 'col_count': cols,
            'columns': [{'index': c, 'header': header[c] if c < len(header) else '',
                         'kind': _column_kind(body_values[c])} for c in range(cols)],
        })
    return out


# ── Temporary storage of a parsed upload ──────────────────────────────────

def _import_dir() -> str:
    path = os.path.join(current_app.instance_path, 'sheet_imports')
    os.makedirs(path, exist_ok=True)
    return path


def _cleanup(path: str):
    cutoff = time.time() - TOKEN_TTL_SECONDS
    for name in os.listdir(path):
        full = os.path.join(path, name)
        try:
            if os.path.getmtime(full) < cutoff:
                os.remove(full)
        except OSError:
            pass


def store(user_id: int, filename: str, sheets: list) -> str:
    path = _import_dir()
    _cleanup(path)
    token = secrets.token_urlsafe(24)
    with open(os.path.join(path, f'{token}.json'), 'w', encoding='utf-8') as fh:
        json.dump({'user_id': user_id, 'filename': filename, 'sheets': sheets}, fh)
    return token


def load(token: str, user_id: int) -> dict:
    if not isinstance(token, str) or not _TOKEN_RE.match(token):
        raise ImportError_('This import has expired. Upload the file again.')
    full = os.path.join(_import_dir(), f'{token}.json')
    if not os.path.exists(full) or os.path.getmtime(full) < time.time() - TOKEN_TTL_SECONDS:
        raise ImportError_('This import has expired. Upload the file again.')
    with open(full, encoding='utf-8') as fh:
        data = json.load(fh)
    if data.get('user_id') != user_id:
        raise ImportError_('This import has expired. Upload the file again.')
    return data


def discard(token: str):
    if isinstance(token, str) and _TOKEN_RE.match(token):
        try:
            os.remove(os.path.join(_import_dir(), f'{token}.json'))
        except OSError:
            pass


# ── Export ────────────────────────────────────────────────────────────────

_NUM_RE = re.compile(r'^-?\d+(\.\d+)?([eE][-+]?\d+)?$')
_DATE_RE = re.compile(r'^(\d{4})-(\d{2})-(\d{2})(?:[ T](\d{2}):(\d{2})(?::(\d{2}))?)?$')


def _excel_value(raw: Optional[str]):
    if raw is None:
        return None
    if raw.startswith("'"):
        return raw[1:]
    if raw.startswith('='):
        return raw
    if raw.upper() in ('TRUE', 'FALSE'):
        return raw.upper() == 'TRUE'
    plain = raw.replace(',', '') if re.match(r'^-?\d{1,3}(,\d{3})+(\.\d+)?$', raw) else raw
    if _NUM_RE.match(plain):
        n = float(plain)
        return int(n) if n.is_integer() and '.' not in plain and 'e' not in plain.lower() else n
    if plain.endswith('%') and _NUM_RE.match(plain[:-1]):
        return float(plain[:-1]) / 100
    m = _DATE_RE.match(raw)
    if m:
        try:
            if m.group(4):
                return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)),
                                int(m.group(4)), int(m.group(5)), int(m.group(6) or 0))
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return raw
    return raw


def _excel_number_format(nf: Optional[dict], value) -> Optional[str]:
    if not nf:
        if isinstance(value, datetime):
            return 'yyyy-mm-dd' if value.hour == value.minute == 0 else 'yyyy-mm-dd hh:mm'
        return None
    t, d = nf.get('t'), nf.get('d')
    decimals = ('.' + '0' * d) if d else ''
    if t == 'number':
        return '#,##0' + decimals
    if t == 'decimal':
        return '#,##0' + ('.' + '0' * (d if d is not None else 2))
    if t == 'currency':
        cur = (nf.get('cur') or 'Rs').replace('"', '')
        return f'"{cur} "#,##0' + ('.' + '0' * d if d else '')
    if t == 'percent':
        return '0' + decimals + '%'
    if t == 'date':
        return 'yyyy-mm-dd'
    if t == 'time':
        return 'hh:mm:ss'
    if t == 'datetime':
        return 'yyyy-mm-dd hh:mm'
    if t == 'text':
        return '@'
    return None


def export_xlsx(spreadsheet, sheets_with_cells) -> bytes:
    """sheets_with_cells: [(SpreadsheetSheet, [SpreadsheetCell])]."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used_titles = set()
    for sheet, cells in sheets_with_cells:
        title = re.sub(r'[\[\]\*\?/\\:]', '_', sheet.name)[:31] or 'Sheet'
        base, n = title, 2
        while title.lower() in used_titles:
            title = f'{base[:27]} ({n})'
            n += 1
        used_titles.add(title.lower())
        ws = wb.create_sheet(title)
        props = V.load(sheet.props, {}) or {}
        for cell in cells:
            fmt = V.load(cell.fmt, {}) or {}
            value = _excel_value(cell.value)
            if value is None and not fmt:
                continue
            xc = ws.cell(row=cell.row + 1, column=cell.col + 1)
            if value is not None:
                xc.value = value
            if fmt.get('nf', {}).get('t') == 'text' and value is not None and not isinstance(value, str):
                xc.value = cell.value
            num_fmt = _excel_number_format(fmt.get('nf'), value)
            if num_fmt:
                xc.number_format = num_fmt
            if any(k in fmt for k in ('b', 'i', 'u', 's', 'fs', 'ff', 'c')):
                xc.font = Font(bold=fmt.get('b', False), italic=fmt.get('i', False),
                               underline='single' if fmt.get('u') else None, strike=fmt.get('s', False),
                               size=fmt.get('fs', 11), name=fmt.get('ff', 'Calibri'),
                               color=('FF' + fmt['c'][1:7].upper()) if fmt.get('c') and len(fmt['c']) >= 7 else None)
            if fmt.get('bg') and len(fmt['bg']) >= 7:
                xc.fill = PatternFill('solid', fgColor='FF' + fmt['bg'][1:7].upper())
            if any(k in fmt for k in ('ha', 'va', 'wrap')):
                xc.alignment = Alignment(horizontal=fmt.get('ha'),
                                         vertical={'middle': 'center'}.get(fmt.get('va'), fmt.get('va')),
                                         wrap_text=fmt.get('wrap') == 'wrap')
            if any(k in fmt for k in ('bt', 'bb', 'bl', 'br')):
                def side(spec):
                    if not spec:
                        return Side()
                    style = {'dashed': 'dashed', 'dotted': 'dotted', 'double': 'double'}.get(
                        spec.get('s'), 'medium' if spec.get('w', 1) >= 2 else 'thin')
                    return Side(style=style, color='FF' + spec.get('c', '#000000')[1:7].upper())
                xc.border = Border(top=side(fmt.get('bt')), bottom=side(fmt.get('bb')),
                                   left=side(fmt.get('bl')), right=side(fmt.get('br')))
        for col, width in (props.get('cw') or {}).items():
            ws.column_dimensions[get_column_letter(int(col) + 1)].width = max(1, (int(width) - 5) / 7)
        for row, height in (props.get('rh') or {}).items():
            ws.row_dimensions[int(row) + 1].height = int(height) * 0.75
        for col in props.get('hc') or []:
            ws.column_dimensions[get_column_letter(int(col) + 1)].hidden = True
        for row in props.get('hr') or []:
            ws.row_dimensions[int(row) + 1].hidden = True
        for m in props.get('merges') or []:
            ws.merge_cells(start_row=m['r1'] + 1, start_column=m['c1'] + 1,
                           end_row=m['r2'] + 1, end_column=m['c2'] + 1)
        if sheet.frozen_rows or sheet.frozen_cols:
            ws.freeze_panes = ws.cell(row=sheet.frozen_rows + 1, column=sheet.frozen_cols + 1)
        if sheet.color and len(sheet.color) >= 7:
            ws.sheet_properties.tabColor = sheet.color[1:7].upper()
        if sheet.is_hidden:
            ws.sheet_state = 'hidden'
    if not wb.worksheets:
        wb.create_sheet('Sheet1')
    if all(ws.sheet_state == 'hidden' for ws in wb.worksheets):
        wb.worksheets[0].sheet_state = 'visible'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
